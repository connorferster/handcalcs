from __future__ import annotations

import pathlib
from typing import Any, ClassVar, Optional

from handcalcs.renderers.base import BaseRenderContext, BaseRenderer, RenderContext
from . import omml
from .document import (
    CalcGrid,
    Code,
    DocumentWriter,
    Heading,
    MathGroup,
    ParamsTable,
    Prose,
    Row,
    validate_layout,
)


class WordRenderer(BaseRenderer):
    """
    Renders a HandCalcs sequence as a Word document (.docx) whose calculations
    are native, editable Word equations (Office Math Markup, OMML).

    Like the LaTeX renderer, the node handlers return math as strings, here OMML
    fragments (see ``omml``), and the base renderer's line and block machinery
    works on them unchanged. ``complete()`` writes the render tree as a
    python-docx ``Document``; python-docx (``pip install handcalcs[word]``) is
    imported only then.

    Options:

    - ``filepath``: if given, ``complete()`` also saves the document there.
    - ``layout``: how a run of calc lines is set out. ``"aligned"`` (the
      default) is one equation array per run, aligned on the first ``=`` of
      each row, as in the LaTeX renderer. ``"lines"`` makes each row an
      equation of its own with no alignment; use it for Google Docs and
      LibreOffice, which don't show Word's equation arrays. A ``# hc: layout =
      "lines"`` comment changes it for the calcs that follow.
    - ``template``: a .docx to start each new document from (a company
      template: its styles, page setup and any content it already has).
    - ``document``: a python-docx ``Document`` to append to instead, so several
      calcs build one report. Give a template or a document, not both.

    The renderer adds two styles when the document doesn't define them:
    ``HC Math`` for equation paragraphs and ``HC Code`` for imports. Define
    them in a template to restyle both. Equations take their paragraph's font
    size.

    Examples::

        HandCalcs(WordRenderer(filepath="calc.docx"))(source)

        report = docx.Document("company.docx")
        calc = HandCalcs(WordRenderer(document=report))
        calc(beam_source)
        calc(column_source)
        report.save("report.docx")
    """

    name = "word"

    # A continuation row (a source line continuation) starts with an em space.
    CONTINUATION_INDENT: ClassVar[str] = omml.run(omml.EM_SPACE)
    # An inline comment starts with COMMENT_OPEN (see ``render_inline_comment``).
    COMMENT_OPEN: ClassVar[str] = omml.run(omml.EM_SPACE)
    # Collections longer than MAX_ITEMS show only their first and last
    # EDGE_ITEMS items (an equation row cannot break).
    MAX_ITEMS: ClassVar[int] = 10
    EDGE_ITEMS: ClassVar[int] = 3
    # Keywords are upright words, like ``and``/``or``.
    IF_EXP_KEYWORDS: ClassVar[tuple[str, str]] = (
        omml.upright(" if "),
        omml.upright(" else "),
    )
    COMPREHENSION_KEYWORDS: ClassVar[tuple[str, str]] = (
        omml.upright(" for "),
        omml.upright(" in "),
    )

    def __init__(
        self,
        filepath: Optional[str | pathlib.Path] = None,
        layout: str = "aligned",
        template: Optional[str | pathlib.Path] = None,
        document: Any = None,
    ) -> None:
        super().__init__()
        if template is not None and document is not None:
            raise ValueError("Give a template or a document, not both.")
        self.filepath = pathlib.Path(filepath) if filepath is not None else None
        self.layout = validate_layout(layout)
        self.template = pathlib.Path(template) if template is not None else None
        if self.template is not None and not self.template.is_file():
            raise FileNotFoundError(f"No Word template at {self.template}")
        self.document = document

    def __repr__(self) -> str:
        return f"WordRenderer(layout={self.layout!r})"

    def create_context(self, **kwargs) -> BaseRenderContext:
        # Word spaces math itself, so the base handlers join components with no
        # space; brackets are the two halves of a growing ``m:d`` delimiter.
        return BaseRenderContext(
            global_context=RenderContext(
                space="",
                newline="\n",
                indent="",
                lpar=omml.delim_open(),
                rpar=omml.DELIM_CLOSE,
                **kwargs,
            ),
            line_context=RenderContext.sparse(),
        )

    def leading_sign(self, rendered: str) -> bool:
        return omml.plain_text(rendered).lstrip().startswith(("-", "+", "−"))

    def complete(self, tree: list, base_context: BaseRenderContext) -> Any:
        """
        Write the render tree as a Word document and return it; save it if
        ``filepath`` is set. The calc is appended to ``document`` if one was
        given, or else to a new document (from ``template``, if given).
        """
        if self.document is not None:
            document = self.document
        else:
            from docx import Document

            document = Document(self.template) if self.template else Document()
        writer = DocumentWriter(document)
        self.write_items(writer, tree, 0, base_context)
        if self.filepath is not None:
            document.save(self.filepath)
        return document

    # --- Writing the tree ----------------------------------------------------

    def write_items(
        self, writer: DocumentWriter, items: list, depth: int, base_context
    ) -> None:
        for item in items:
            self.write_item(writer, item, depth, base_context)

    def write_item(
        self, writer: DocumentWriter, item, depth: int, base_context
    ) -> None:
        """
        Word counterpart of ``BaseRenderer.join_item``. Master-list shapes:

        - a ``Heading``, ``Code`` or ``ParamsTable`` item -> its paragraph or table;
        - a string -> a text paragraph; a one-shot line break -> an empty one;
        - a headerless block whose body opens with a ``MathGroup`` (a
          CalcsBlock) -> its equations (see ``DocumentWriter.math_group``);
        - another headerless block (a CommentsBlock) -> its body at this level;
        - ``[header, body]`` -> the header paragraph, then the body indented
          one level;
        - a calc row (a list of strings) or a ``CalcGrid`` on its own -> its
          equations; a list of plain strings -> a text paragraph.
        """
        if not item:
            return
        if isinstance(item, Heading):
            writer.heading(item)
        elif isinstance(item, Code):
            writer.code(item, depth)
        elif isinstance(item, ParamsTable):
            writer.params_table(item, depth)
        elif isinstance(item, str):
            if item.strip():
                writer.prose(item, depth)
            else:
                writer.line_break(depth)
        elif isinstance(item, CalcGrid) or not isinstance(item[-1], list):
            rows = self.math_rows([item], base_context)
            if rows:
                writer.math_group(rows, self.layout, depth)
        else:
            header, body = item[0], item[-1]
            if body and isinstance(body[0], MathGroup):
                rows = self.math_rows(body[1:], base_context)
                if rows:
                    writer.math_group(rows, body[0].layout, depth)
            elif isinstance(header, str) and header.strip():
                writer.prose(header, depth)
                self.write_items(writer, body, depth + 1, base_context)
            else:
                self.write_items(writer, body, depth, base_context)

    def math_rows(self, lines: list, base_context) -> list[Row | Prose]:
        """The equation rows of a CalcsBlock's lines (see ``join_math_row``)."""
        rows: list[Row | Prose] = []
        for line in lines:
            if isinstance(line, CalcGrid):
                rows.extend(self.grid_rows(line))
            elif isinstance(line, str):
                if line.strip():
                    rows.append(Prose(line))
            elif line:
                rows.extend(self.join_math_row(line, base_context))
        return rows

    def join_math_row(self, parts: list, base_context) -> list[Row | Prose]:
        """
        One calc line's components as an equation row, split before its first
        assignment symbol (the alignment point). A line with no math (a
        docstring) is prose; an ignored line is no row.

        As in the LaTeX renderer, the inline comment of a line that shows
        substitution columns (more than ``target = value``) moves to its own
        row, at the alignment point, so it does not lengthen a long row.
        """
        context = base_context.current
        parts = list(parts)
        if not parts or parts[0] == "":
            return []
        if not any("<m:" in part for part in parts):
            return [Prose(" ".join(parts))]
        comment_row = []
        if parts[-1].startswith(self.COMMENT_OPEN) and parts.count(context.assign) > 1:
            comment_row = [Row("", parts.pop())]
        if context.assign in parts:
            idx = parts.index(context.assign)
        else:
            idx = 0
        row = Row("".join(parts[:idx]), "".join(parts[idx:]))
        return [row, *comment_row]

    def grid_rows(self, grid: CalcGrid) -> list[Row]:
        """
        The rows of a multiline calc, aligned on ``eq``. A row with an empty
        ``eq`` continues the row above and starts with ``CONTINUATION_INDENT``.
        """
        rows = []
        for ident, eq, val in grid.rows:
            if not eq:
                val = f"{self.CONTINUATION_INDENT}{val}"
            rows.append(Row(ident, f"{eq}{val}"))
        return rows

    def format_param_grid(self, rows: list, base_context) -> ParamsTable:
        return ParamsTable([[cell and list(cell) for cell in row] for row in rows])

    def format_calc_grid(self, rows: list, base_context) -> CalcGrid:
        return CalcGrid([list(row) for row in rows])
