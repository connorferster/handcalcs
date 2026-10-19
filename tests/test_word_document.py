"""
Tests for the Word renderer's document layout (#222, step 4): calc blocks as
equation arrays (or one equation per row), params tables, prose, headings and
nested blocks. Each test renders a source, then reads the python-docx
``Document``'s body back.
"""

import pytest

docx = pytest.importorskip("docx")

from docx.oxml.ns import qn

from handcalcs import HandCalcs
from handcalcs.renderers.word import WordRenderer, omml

M = f"{{{omml.M_NS}}}"


def render(source: str, **options):
    return HandCalcs(renderer=WordRenderer(**options))(source)


def math_text(el) -> str:
    return "".join(t.text or "" for t in el.iter(f"{M}t"))


def eq_rows(paragraph) -> list[str]:
    """The text of each row of a paragraph's equation array."""
    return [math_text(row) for row in paragraph.findall(f".//{M}eqArr/{M}e")]


def indent(paragraph) -> int:
    ind = paragraph.find(f"{qn('w:pPr')}/{qn('w:ind')}")
    return int(ind.get(qn("w:left"))) if ind is not None else 0


def body(document) -> list:
    """The body's paragraphs and tables (without the section properties)."""
    return [el for el in document.element.body if el.tag != qn("w:sectPr")]


def math_paragraphs(document) -> list:
    return [el for el in body(document) if el.find(f"{M}oMathPara") is not None]


# --- calc blocks: the "aligned" layout ---------------------------------------


def test_complete_returns_a_document():
    document = render("a = 2\n")
    assert isinstance(document, docx.document.Document)


def test_calc_block_is_one_equation_array():
    document = render("a = 2\nb = 3\nc = a * b\n")
    (paragraph,) = math_paragraphs(document)
    assert eq_rows(paragraph) == ["a&=2", "b&=3", "c&=a⋅b=2⋅3=6"]
    jc = paragraph.find(f"{M}oMathPara/{M}oMathParaPr/{M}jc")
    assert jc.get(f"{M}val") == "left"


def test_alignment_point_is_its_own_run():
    document = render("a = 2\n")
    (paragraph,) = math_paragraphs(document)
    row = paragraph.find(f".//{M}eqArr/{M}e")
    runs = [math_text(run) for run in row.findall(f"{M}r")]
    assert runs == ["a", "&", "=", "2"]


def test_comment_moves_to_its_own_row_with_substitutions():
    document = render("a = 2\nb = a * 3  # note\n")
    (paragraph,) = math_paragraphs(document)
    assert eq_rows(paragraph) == ["a&=2", "b&=a⋅3=2⋅3=6", "&\u2003(note)"]


def test_comment_stays_on_a_short_row():
    document = render("a = 2  # note\n")
    (paragraph,) = math_paragraphs(document)
    assert eq_rows(paragraph) == ["a&=2\u2003(note)"]


def test_multiline_calc_rows():
    document = render("a = 2\n# hc: -m\nb = a * 3\n")
    rows = [row for p in math_paragraphs(document) for row in eq_rows(p)]
    assert rows[-3:] == ["b&=a⋅3", "&=2⋅3", "&=6"]


def test_continuation_rows_start_with_an_em_space():
    document = render("a = 2\nb = a + \\\n    a * 3\n")
    rows = [row for p in math_paragraphs(document) for row in eq_rows(p)]
    assert rows[1:] == ["b&=a+", "&\u2003a⋅3", "&=2+", "&\u20032⋅3", "&=8"]


def test_ignored_lines_make_no_rows():
    document = render("a = 2  # hc: -i\nb = 3\n")
    (paragraph,) = math_paragraphs(document)
    assert eq_rows(paragraph) == ["b&=3"]


# --- the "lines" layout --------------------------------------------------------


def test_lines_layout_is_one_equation_per_row():
    document = render("a = 2\n# hc: -m\nb = a * 3\n", layout="lines")
    paragraphs = math_paragraphs(document)
    assert [math_text(p) for p in paragraphs] == ["a=2", "b=a⋅3", "=2⋅3", "=6"]
    assert [indent(p) for p in paragraphs] == [0, 0, 360, 360]
    assert not document.element.body.findall(f".//{M}eqArr")


def test_layout_comment_applies_to_the_calcs_after_it():
    source = 'a = 2\n# hc: layout = "lines"\nb = 3\nc = 4\n'
    document = render(source)
    arrays = document.element.body.findall(f".//{M}eqArr")
    assert [math_text(arr) for arr in arrays] == ["a&=2"]
    assert [math_text(p) for p in math_paragraphs(document)][1:] == ["b=3", "c=4"]


def test_unknown_layout_is_an_error():
    with pytest.raises(ValueError, match="layout"):
        WordRenderer(layout="table")
    with pytest.raises(ValueError, match="layout"):
        render('# hc: layout = "table"\na = 2\n')


# --- params, prose and blocks ----------------------------------------------------


def test_params_grid_is_a_borderless_table():
    document = render("b = 300; d = 500; f_c = 30; t = 2\n")
    (table,) = document.tables
    assert len(table.rows) == 2 and len(table.columns) == 3
    cells = [math_text(cell._tc) for row in table.rows for cell in row.cells]
    assert cells == ["b=300", "d=500", "fc=30", "t=2", "", ""]
    assert table.style is None or "Grid" not in table.style.name


def test_heading_comment_and_import():
    document = render("## Loads\n# The {{ 'dead' }} load\n# in kN\nimport math\n")
    heading, prose, code = document.paragraphs
    assert (heading.style.name, heading.text) == ("Heading 1", "Loads")
    assert prose.text == "The dead load in kN"
    assert code.text == "import math"
    assert code.runs[0].font.name == "Consolas"


def test_line_break_is_an_empty_paragraph():
    document = render("a = 2\n# hc: -b\nb = 3\n")
    assert [p.text for p in document.paragraphs] == ["", "", ""]
    assert len(math_paragraphs(document)) == 2


def test_if_block_header_and_indented_body():
    source = "a = 2\nif a > 1:\n    b = 3\n    if b > 2:\n        c = 4\n"
    document = render(source)
    headers = [p for p in body(document) if "Since" in "".join(p.itertext())]
    assert len(headers) == 2
    first = headers[0]
    assert first.find(f"{M}oMath") is not None
    assert "".join(t.text for t in first.iter(qn("w:t"))) == "Since  is True:"
    assert math_text(first) == "a>1→2>1"
    assert [indent(h) for h in headers] == [0, 720]
    assert [(eq_rows(p), indent(p)) for p in math_paragraphs(document)] == [
        (["a&=2"], 0),
        (["b&=3"], 720),
        (["c&=4"], 1440),
    ]


def test_nested_params_table_is_indented():
    document = render("a = 2\nif a > 1:\n    b = 3; c = 4\n")
    (table,) = document.tables
    ind = table._tbl.tblPr.find(qn("w:tblInd"))
    assert ind.get(qn("w:w")) == "720"


def test_for_block_header():
    document = render("total = 0\nfor i in [1, 2]:\n    total = total + i\n")
    header = next(p for p in body(document) if "Iterating" in "".join(p.itertext()))
    assert len(header.findall(f"{M}oMath")) == 2


def test_filepath_saves_the_document(tmp_path):
    path = tmp_path / "calc.docx"
    render("## Check\na = 2\nb = a * 3\n", filepath=path)
    reopened = docx.Document(str(path))
    assert reopened.paragraphs[0].style.name == "Heading 1"
    (paragraph,) = math_paragraphs(reopened)
    assert eq_rows(paragraph) == ["a&=2", "b&=a⋅3=2⋅3=6"]


@pytest.mark.parametrize("layout", ["aligned", "lines"])
def test_demo_renders(layout, tmp_path):
    path = tmp_path / "demo.docx"
    HandCalcs(renderer=WordRenderer(filepath=path, layout=layout)).demo()
    reopened = docx.Document(str(path))
    assert math_paragraphs(reopened)
    assert reopened.tables
