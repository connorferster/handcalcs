"""
A renderer-neutral test suite of handcalcs' features, for every renderer: the
built-in ones and those in other packages (such as ym-handcalcs).

Each ``FeatureCase`` is a short calculation that uses one feature, and names the
variables the output must show and the ones it must not. A renderer passes a case
when it renders the source without an error and the text of its output has every
``shown`` name and none of the ``hidden`` ones. The names are plain words (fruit)
that no renderer turns into symbols, so the same check works for every output.

A renderer package runs the suite from its own tests, with a function that turns
its output into text (``output_text`` already knows strings, Word documents and
Excel workbooks)::

    import pytest
    from handcalcs.testing import FEATURE_CASES, check_case

    @pytest.mark.parametrize("case", FEATURE_CASES, ids=lambda case: case.id)
    def test_feature(case):
        check_case(MyRenderer(), case, text_of=my_output_text)

A new feature adds a case here, so every renderer is tested on it from the next
handcalcs release on.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any, Callable, Mapping

from handcalcs.config import get_config, reset_option, set_option


@dataclass(frozen=True)
class FeatureCase:
    """
    One feature, as a calculation: ``source`` is rendered with ``options`` set
    (``set_option``) and its output must contain every name in ``shown`` and none
    in ``hidden``.
    """

    id: str
    source: str
    shown: tuple[str, ...]
    hidden: tuple[str, ...] = ()
    options: Mapping[str, Any] = field(default_factory=dict)


def _case(id: str, source: str, shown: str, hidden: str = "", **options) -> FeatureCase:
    return FeatureCase(id, source, tuple(shown.split()), tuple(hidden.split()), options)


FEATURE_CASES: tuple[FeatureCase, ...] = (
    # --- Lines ---
    _case("param", "apple = 3\n", "apple"),
    _case("calc", "apple = 3\nberry = apple * 2 + 1\n", "apple berry"),
    _case("expression", "apple = 3\napple * 2 + 1\n", "apple"),
    _case(
        "heading-and-comment",
        "## Grape\n# A lemon note\napple = 3\n",
        "Grape lemon apple",
    ),
    _case("inline-comment", "apple = 3  # lemon note\n", "apple lemon"),
    _case("import", "from math import sqrt\napple = sqrt(9)\n", "apple"),
    _case("augmented", "apple = 1\napple += 2\n", "apple"),
    _case("collection", "apple = [1, 2, 3]\nberry = sum(apple)\n", "apple berry"),
    _case("continuation", "apple = 2\nberry = (apple\n    + 3)\n", "apple berry"),
    _case(
        "if-block",
        "apple = 3\nif apple > 2:\n    berry = 1\nelse:\n    cherry = 2\n",
        "apple berry",
    ),
    _case(
        "for-block",
        "melon = 0\nfor plum in [1, 2]:\n    melon = melon + plum\n",
        "melon",
    ),
    # --- Flags ---
    _case("format-code", "apple = 1 / 3  # hc: -f .2f\n", "apple"),
    _case(
        "format-code-region",
        "# hc: -f .3e\napple = 1 / 3\nberry = apple * 2\n",
        "apple berry",
    ),
    _case(
        "multiline",
        "apple = 2\nberry = apple * 3 + apple / 4  # hc: -m\n",
        "apple berry",
    ),
    _case("param-line", "apple = 2\nberry = apple * 3  # hc: -p\n", "apple berry"),
    _case(
        "grid",
        "# hc: -g -c 2\napple = 1  # lemon\nberry = 2\ncherry = 3\n",
        "apple berry cherry lemon",
    ),
    _case(
        "grid-semicolons", "apple = 1; berry = 2; cherry = 3\n", "apple berry cherry"
    ),
    _case("line-break", "apple = 1\n# hc: -b\nberry = 2\n", "apple berry"),
    _case(
        "check",
        "apple = 3\nberry = apple > 2  # hc: --check\ncherry = apple < 2  # hc: --check\n",
        "apple berry cherry",
    ),
    _case(
        "pagebreak",
        "apple = 1\n# hc: --pagebreak\nberry = 2  # hc: --pagebreak\ncherry = 3\n",
        "apple berry cherry",
    ),
    _case("winner-only", "apple = 3  # hc: -w\n", "apple"),
    # --- Assignment form ---
    _case(
        "mode", '# hc: mode = "ass res"\napple = 2\nberry = apple * 3\n', "apple berry"
    ),
    _case(
        "assign", '# hc: assign = ":="\napple = 2\nberry = apple * 3\n', "apple berry"
    ),
    _case(
        "subscripts-nested",
        '# hc: subscripts = "nested"\napple = 2\nberry = apple * 3\n',
        "apple berry",
    ),
    _case(
        "symbols",
        '# hc: symbols = {"cherry": "C_{x}"}\napple = 2\ncherry = apple * 3\n',
        "apple",
        "cherry",
    ),
    _case(
        "greek-prefixes-off",
        "# hc: greek_prefixes = False\napple = 2\nberry = apple\n",
        "apple berry",
    ),
    # --- Showing and hiding ---
    _case(
        "ignore-inline",
        "apple = 1\nberry = 2  # hc: -i\ncherry = 3\n",
        "apple cherry",
        "berry",
    ),
    _case(
        "ignore-region",
        "apple = 1\n# hc: -i\n## Grape\n# A lemon note\nberry = 2\nif berry > 1:\n"
        "    cherry = 3\n# hc: -g\nkiwi = 1\nmango = 2\n# hc: -s\nmelon = 4\n",
        "apple melon",
        "Grape lemon berry cherry kiwi mango",
    ),
    _case(
        "ignore-assignment",
        "apple = 1\n# hc: ignore = True\nberry = 2\n# hc: ignore = False\ncherry = 3\n",
        "apple cherry",
        "berry",
    ),
    _case(
        "show-inline",
        "# hc: -i\napple = 1\nberry = 2  # hc: -s\ncherry = 3\n",
        "berry",
        "apple cherry",
    ),
    _case(
        "hidden-grid-line",
        "apple = 1; berry = 2  # hc: -i\ncherry = 3\n",
        "apple cherry",
        "berry",
    ),
    _case(
        "hidden-block",
        "apple = 3\n# hc: -i\nif apple > 2:\n    berry = 1\n    # hc: -s\n    cherry = 2\nmelon = 4\n",
        "apple melon",
        "berry cherry",
    ),
    _case(
        "show-condition-false",
        "apple = 0\n# hc: -s apple > 0\n## Grape\nberry = apple * 2\n# hc: -s\ncherry = 3\n",
        "apple cherry",
        "Grape berry",
    ),
    _case(
        "show-condition-true",
        "apple = 1\n# hc: -s apple > 0\n## Grape\nberry = apple * 2\n# hc: -s\ncherry = 3\n",
        "apple Grape berry cherry",
    ),
    _case(
        "ignore-condition-inline",
        "apple = 2\nberry = apple - 2  # hc: -i berry == 0\ncherry = apple + 1  # hc: -i cherry == 0\n",
        "apple cherry",
        "berry",
    ),
    _case(
        "hidden-by-default",
        "## Grape\napple = 1\nberry = 2  # hc: -s\n# hc: -s\ncherry = 3\n# hc: -i\nmelon = 4\n",
        "berry cherry",
        "Grape apple melon",
        hidden_by_default=True,
    ),
)

_MISSING = object()


def render_case(renderer: Any, case: FeatureCase) -> Any:
    """Render ``case.source`` with ``renderer``, with ``case.options`` set for the render."""
    from handcalcs import HandCalcs

    before = {name: get_config().get(name, _MISSING) for name in case.options}
    try:
        for name, value in case.options.items():
            set_option(name, value)
        return HandCalcs(renderer=renderer)(case.source)
    finally:
        for name, value in before.items():
            if value is _MISSING:
                reset_option(name)
            else:
                set_option(name, value)


def output_text(rendered: Any) -> str:
    """
    The visible text of a built-in renderer's output: a string as it is, the text of
    a Word document (body and equations), or the cells of an Excel workbook's shown
    rows (a hidden row is left out).
    """
    if isinstance(rendered, str):
        return rendered
    element = getattr(rendered, "element", None)
    if element is not None:  # a python-docx Document
        return " ".join(element.body.itertext())
    if hasattr(rendered, "to_bytes"):  # an ExcelWorkbook
        import io

        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(rendered.to_bytes())).active
        return " ".join(
            str(cell.value)
            for row in sheet.iter_rows()
            if row and not sheet.row_dimensions[row[0].row].hidden
            for cell in row
            if cell.value is not None
        )
    raise TypeError(
        f"output_text() does not know {type(rendered).__name__}; pass check_case() a "
        "'text_of' function that returns the text of your renderer's output."
    )


def check_text(case: FeatureCase, text: str) -> None:
    """Raise ``AssertionError`` unless ``text`` shows every ``shown`` name and no ``hidden`` one."""
    found = (
        lambda name: re.search(rf"(?<![A-Za-z]){name}(?![A-Za-z])", text) is not None
    )
    missing = [name for name in case.shown if not found(name)]
    unexpected = [name for name in case.hidden if found(name)]
    if missing or unexpected:
        raise AssertionError(
            f"feature case {case.id!r}: missing {missing}, should be hidden {unexpected}\n"
            f"source:\n{case.source}\noutput text:\n{text}"
        )


def check_case(
    renderer: Any, case: FeatureCase, text_of: Callable[[Any], str] = output_text
) -> None:
    """Render ``case`` with ``renderer`` and check its text (see ``check_text``)."""
    check_text(case, text_of(render_case(renderer, case)))
