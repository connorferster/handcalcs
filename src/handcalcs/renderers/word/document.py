"""
The Word document side of the renderer: the items the handlers put in the
render tree besides plain strings and calc rows, and ``DocumentWriter``, which
writes the tree's items as paragraphs, equations and tables of a python-docx
``Document``.

python-docx (and lxml) are imported only here and in ``omml.to_element``, when a
document is built.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from typing import Any

from . import omml

LAYOUTS = ("aligned", "lines")

# Twips (1/1440 inch) per block level, and the extra indent of a row with no
# left-hand side in the "lines" layout (a continuation or a numeric row).
INDENT = 720
ROW_INDENT = 360

CODE_FONT = "Consolas"


@dataclass
class Heading:
    text: str
    level: int


@dataclass
class Code:
    """A line of Python source (an import), set in a monospace font."""

    text: str


@dataclass
class ParamsTable:
    """A params grid: rows of cells, each ``[id, eq, value, comment]`` or ``None``."""

    rows: list


@dataclass
class CalcGrid:
    """A multiline calc (``-m`` or a line continuation): ``[id, eq, value]`` rows."""

    rows: list


@dataclass
class MathGroup:
    """Opens the body of a CalcsBlock: its rows form one group of equations."""

    layout: str = "aligned"


@dataclass
class Row:
    """One equation row: the left-hand side and the rest from the ``=`` on."""

    lhs: str
    rhs: str


@dataclass
class Prose:
    """A run of text paragraphs inside a math group (e.g. a docstring line)."""

    text: str


def validate_layout(layout: Any) -> str:
    if layout not in LAYOUTS:
        raise ValueError(f"layout must be one of {LAYOUTS}, not {layout!r}")
    return layout


_INLINE_MATH = re.compile(r"(<m:oMath>.*?</m:oMath>)", re.DOTALL)


def inline_math(fragment: str) -> str:
    """Mark an OMML fragment as inline math inside a line of prose."""
    return omml.omath(fragment)


@dataclass
class DocumentWriter:
    """Appends paragraphs, equations and tables to a python-docx ``Document``."""

    document: Any

    def paragraph(self, depth: int = 0, style: str | None = None):
        from docx.shared import Twips

        paragraph = self.document.add_paragraph(style=style)
        if depth:
            paragraph.paragraph_format.left_indent = Twips(INDENT * depth)
        return paragraph

    def heading(self, item: Heading) -> None:
        self.document.add_heading(item.text, level=min(max(item.level, 1), 9))

    def prose(self, text: str, depth: int = 0) -> None:
        """
        A text paragraph. ``text`` may hold inline equations (``inline_math``),
        which become ``m:oMath`` elements between the text runs.
        """
        paragraph = self.paragraph(depth)
        for piece in _INLINE_MATH.split(text):
            if not piece:
                continue
            if piece.startswith("<m:oMath>"):
                inner = piece[len("<m:oMath>") : -len("</m:oMath>")]
                paragraph._p.append(omml.to_element(inner))
            else:
                paragraph.add_run(piece)

    def code(self, item: Code, depth: int = 0) -> None:
        run = self.paragraph(depth).add_run(item.text)
        run.font.name = CODE_FONT

    def line_break(self, depth: int = 0) -> None:
        self.paragraph(depth)

    def display_math(self, fragment: str, depth: int = 0, extra: int = 0) -> None:
        """A paragraph holding one flush-left display equation (``m:oMathPara``)."""
        from docx.oxml import parse_xml
        from docx.shared import Twips

        paragraph = self.paragraph(depth)
        if extra:
            paragraph.paragraph_format.left_indent = Twips(INDENT * depth + extra)
        para = parse_xml(
            f"<m:oMathPara {omml.NSMAP}>"
            '<m:oMathParaPr><m:jc m:val="left"/></m:oMathParaPr></m:oMathPara>'
        )
        para.append(omml.to_element(fragment))
        paragraph._p.append(para)

    def math_group(self, rows: list[Row | Prose], layout: str, depth: int) -> None:
        """
        A CalcsBlock's equations. In the "aligned" layout, each run of rows is
        one equation array aligned before the first ``=`` (a prose row ends the
        run); in the "lines" layout every row is an equation of its own, and a
        row with no left-hand side is indented a little further.
        """
        pending: list[Row] = []

        def flush() -> None:
            if pending:
                arr = omml.eq_arr(
                    [f"{row.lhs}{omml.ALIGN}{row.rhs}" for row in pending]
                )
                self.display_math(arr, depth)
                pending.clear()

        for row in rows:
            if isinstance(row, Prose):
                flush()
                self.prose(row.text, depth)
            elif layout == "lines":
                self.display_math(
                    f"{row.lhs}{row.rhs}", depth, 0 if row.lhs else ROW_INDENT
                )
            else:
                pending.append(row)
        flush()

    def params_table(self, item: ParamsTable, depth: int = 0) -> None:
        """
        A params grid as a borderless table, one inline equation per cell, so
        the grid keeps its columns. A nested block indents the table.
        """
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import parse_xml
        from docx.oxml.ns import qn

        rows = item.rows
        ncols = max((len(row) for row in rows), default=0)
        if not ncols:
            return
        table = self.document.add_table(rows=len(rows), cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        if depth:
            # tblPr children keep the schema order: ... jc, tblInd, ...
            table._tbl.tblPr.find(qn("w:jc")).addnext(
                parse_xml(
                    f'<w:tblInd {omml.NSMAP} w:w="{INDENT * depth}" w:type="dxa"/>'
                )
            )
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                if cell is None:
                    continue
                paragraph = table.cell(i, j).paragraphs[0]
                paragraph._p.append(omml.to_element("".join(map(str, cell))))
